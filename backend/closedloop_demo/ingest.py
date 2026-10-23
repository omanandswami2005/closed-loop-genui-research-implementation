"""Extract plain text from an uploaded document (PDF, PPTX, TXT, MD) or pasted text."""

from __future__ import annotations

import io
import re
from dataclasses import dataclass

MAX_BYTES = 15 * 1024 * 1024
MAX_CHARS = 60_000  # sent to Gemini; longer material is truncated (and says so)
MIN_CHARS = 200  # less than this is not enough to build topics from

KINDS = {".pdf": "pdf", ".pptx": "pptx", ".txt": "text", ".md": "markdown", ".markdown": "markdown"}


class IngestError(ValueError):
    """A problem the learner can fix (wrong type, empty, scanned PDF...)."""


@dataclass(frozen=True)
class Extracted:
    text: str
    kind: str
    units: int  # pages, slides or lines
    truncated: bool


def _kind(filename: str) -> str:
    name = filename.lower()
    for ext, kind in KINDS.items():
        if name.endswith(ext):
            return kind
    raise IngestError("Unsupported file type. Use PDF, PPTX, TXT or MD, or paste the text.")


def _clean(text: str) -> str:
    text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _pdf(data: bytes) -> tuple[str, int]:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise IngestError("This PDF is password-protected. Remove the password or paste the text.")
        pages = [page.extract_text() or "" for page in reader.pages]
    except (PdfReadError, ValueError, KeyError) as exc:
        if isinstance(exc, IngestError):
            raise
        raise IngestError("This PDF could not be read. Try another file or paste the text.") from exc
    return "\n\n".join(f"[Page {i + 1}]\n{p}" for i, p in enumerate(pages) if p.strip()), len(pages)


def _pptx(data: bytes) -> tuple[str, int]:
    from pptx import Presentation

    try:
        deck = Presentation(io.BytesIO(data))
    except Exception as exc:  # python-pptx raises several unrelated types on bad input
        raise IngestError("This PPTX could not be read. Try another file or paste the text.") from exc
    slides = []
    for i, slide in enumerate(deck.slides):
        lines = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                lines += [p.text for p in shape.text_frame.paragraphs if p.text.strip()]
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    lines.append(" | ".join(cell.text for cell in row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(f"Notes: {notes}")
        if lines:
            slides.append(f"[Slide {i + 1}]\n" + "\n".join(lines))
    return "\n\n".join(slides), len(deck.slides)


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def finish(text: str, kind: str, units: int) -> Extracted:
    text = _clean(text)
    if len(text) < MIN_CHARS:
        if kind == "pdf":
            raise IngestError("No readable text found in this PDF (it may be scanned images). Paste the text instead.")
        raise IngestError(f"Too little text to build topics from (need at least {MIN_CHARS} characters).")
    truncated = len(text) > MAX_CHARS
    return Extracted(text[:MAX_CHARS], kind, units, truncated)


def extract(filename: str, data: bytes) -> Extracted:
    if not data:
        raise IngestError("The file is empty.")
    if len(data) > MAX_BYTES:
        raise IngestError("The file is larger than 15 MB.")
    kind = _kind(filename)
    if kind == "pdf":
        text, units = _pdf(data)
    elif kind == "pptx":
        text, units = _pptx(data)
    else:
        text = _decode(data)
        units = text.count("\n") + 1
    return finish(text, kind, units)


def from_paste(text: str) -> Extracted:
    return finish(text, "paste", text.count("\n") + 1)
